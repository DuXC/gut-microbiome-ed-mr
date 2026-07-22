#!/usr/bin/env python3
"""Machine-check the IJIR v0.3 submission package against frozen v0.3 results."""

from __future__ import annotations

import csv
import hashlib
import json
import math
import re
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image


PACKAGE = Path(__file__).resolve().parents[1]
REPOSITORY = Path(__file__).resolve().parents[3]
RESULTS = REPOSITORY / "05_results" / "v0_3_20260722"
SOURCES = PACKAGE / "01_sources"
SUBMISSION = PACKAGE / "03_submission_files"
TABLES = PACKAGE / "05_tables"
SUPPLEMENT = PACKAGE / "06_supplement"
FIGURES = PACKAGE / "04_figures"
QC = PACKAGE / "07_qc"

checks: list[dict[str, object]] = []


def add_check(name: str, passed: bool, evidence: str) -> None:
    checks.append({"check": name, "passed": bool(passed), "evidence": evidence})


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle))


def close(left: float, right: float, tolerance: float = 1e-12) -> bool:
    return math.isclose(float(left), float(right), rel_tol=tolerance, abs_tol=tolerance)


def docx_text(path: Path) -> str:
    doc = Document(path)
    chunks = [paragraph.text for paragraph in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            chunks.extend(cell.text for cell in row.cells)
    return "\n".join(chunks)


def workbook_text_and_errors(path: Path) -> tuple[str, list[str]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    strings: list[str] = []
    errors: list[str] = []
    error_tokens = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!"}
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                text = str(value)
                strings.append(text)
                if text in error_tokens:
                    errors.append(f"{sheet.title}!{cell.coordinate}={text}")
    return "\n".join(strings), errors


required = [
    SUBMISSION / "IJIR_Main_Manuscript_v0_3.docx",
    SUBMISSION / "IJIR_Title_Page_v0_3.docx",
    SUBMISSION / "IJIR_Cover_Letter_v0_3.docx",
    SUBMISSION / "IJIR_STROBE_MR_Checklist_v0_3.docx",
    TABLES / "IJIR_Table_1_Data_Sources_v0_3.xlsx",
    TABLES / "IJIR_Table_2_Forward_Nominal_Associations_v0_3.xlsx",
    SUPPLEMENT / "IJIR_Supplementary_Data_v0_3.xlsx",
    SUPPLEMENT / "IJIR_Supplementary_Figures_v0_3.pdf",
    FIGURES / "IJIR_Figure_1_Study_Design_v0_3_600dpi.png",
    FIGURES / "IJIR_Figure_2_Multiplicity_v0_3_600dpi.png",
    FIGURES / "IJIR_Figure_3_Forward_Evidence_Profile_v0_3_600dpi.png",
    PACKAGE / "CHANGELOG_v0_3.md",
    PACKAGE / "FINAL_QA_REPORT_v0_3.md",
]
required.extend(sorted((FIGURES / "supplementary").glob("*_600dpi.png")))
add_check(
    "required_files_present",
    all(path.is_file() and path.stat().st_size > 0 for path in required),
    f"{sum(path.is_file() and path.stat().st_size > 0 for path in required)}/{len(required)} present",
)

main_source = (SOURCES / "IJIR_Main_Manuscript_v0_3.md").read_text(encoding="utf-8")
title_source = (SOURCES / "IJIR_Title_Page_v0_3.md").read_text(encoding="utf-8")
cover_source = (SOURCES / "IJIR_Cover_Letter_v0_3.md").read_text(encoding="utf-8")
word_report = (QC / "IJIR_Word_Counts_v0_3.txt").read_text(encoding="utf-8")

abstract_words = int(re.search(r"Abstract words: (\d+)", word_report).group(1))
main_words = int(re.search(r"Main-text words .*: (\d+)", word_report).group(1))
references = int(re.search(r"References: (\d+)", word_report).group(1))
cover_words = len(cover_source.split())
add_check("abstract_word_limit", abstract_words <= 200, f"{abstract_words}/200")
add_check("main_text_word_limit", main_words <= 3000, f"{main_words}/3000")
add_check("reference_limit", references <= 50, f"{references}/50")
add_check("cover_letter_length", 350 <= cover_words <= 450, f"{cover_words} whitespace-delimited words")
add_check("main_display_items", "Main figures: 3" in word_report and "Main tables: 2" in word_report, "3 figures + 2 tables")

core_conclusion = (
    "No association met the predefined multiplicity-controlled criteria for a robust "
    "causal interpretation in either direction."
)
add_check("core_conclusion_preserved", core_conclusion in main_source, "exact sentence present")
add_check(
    "title_avoids_independent_replication",
    "independent replication" not in main_source.splitlines()[0].lower(),
    main_source.splitlines()[0].lstrip("# "),
)
add_check(
    "hunt_evidence_named_correctly",
    "cross-cohort same-SNP exposure validation" in main_source
    and "exploratory exact-label HUNT sensitivity" in main_source
    and "P<1×10⁻⁵" in main_source,
    "same-SNP validation plus exploratory P<1×10⁻⁵ legacy analysis disclosed",
)

docx_paths = required[:4]
xlsx_paths = required[4:7]
all_submission_text = "\n".join(docx_text(path) for path in docx_paths)
workbook_texts: list[str] = []
workbook_errors: list[str] = []
for path in xlsx_paths:
    text, errors = workbook_text_and_errors(path)
    workbook_texts.append(text)
    workbook_errors.extend(f"{path.name}:{entry}" for entry in errors)
all_text = "\n".join([main_source, title_source, cover_source, all_submission_text, *workbook_texts])

placeholder_hits = re.findall(r"(?i)\bTODO\b|\bTBD\b|\[placeholder[^]]*\]", all_text)
absolute_path_hits = re.findall(r"/Volumes/[^\s<]+|file://[^\s<]+", all_text)
comma_thousands = re.findall(r"(?<![\d.])\d{1,3},\d{3}(?!\d)", all_text)
add_check("no_placeholders", not placeholder_hits, f"hits={len(placeholder_hits)}")
add_check("no_local_absolute_paths", not absolute_path_hits, f"hits={len(absolute_path_hits)}")
add_check("space_thousands_style", not comma_thousands, f"comma-thousands hits={len(comma_thousands)}")
add_check("no_workbook_formula_errors", not workbook_errors, f"errors={len(workbook_errors)}")

engineering_terms = [
    "evidence geometry",
    "no-go decision",
    "frozen family",
    "cryptographic receipts",
    "stopping rule",
]
engineering_hits = [term for term in engineering_terms if term in main_source.lower()]
add_check("engineering_terms_removed", not engineering_hits, ", ".join(engineering_hits) or "none")

table1 = load_workbook(TABLES / "IJIR_Table_1_Data_Sources_v0_3.xlsx", read_only=True, data_only=True)["Table"]
table1_rows = list(table1.iter_rows(min_row=4, values_only=True))
table2 = load_workbook(TABLES / "IJIR_Table_2_Forward_Nominal_Associations_v0_3.xlsx", read_only=True, data_only=True)["Table"]
table2_rows = list(table2.iter_rows(min_row=4, values_only=True))
add_check("table_1_rows", len(table1_rows) == 9, f"rows={len(table1_rows)}")
add_check("table_2_rows", len(table2_rows) == 7, f"rows={len(table2_rows)}")
add_check(
    "hunt_doi_correct",
    table1_rows[1][8] == "https://doi.org/10.1038/s41588-026-02502-4",
    str(table1_rows[1][8]),
)

multiplicity_rows = read_csv(RESULTS / "multiplicity_sensitivity_trait_results.csv")
primary_by_source = {
    row["source_id"]: row
    for row in multiplicity_rows
    if row["family"] == "forward_primary" and row["p"] and float(row["p"]) < 0.05
}
table2_consistent = len(primary_by_source) == 7
for row in table2_rows:
    source_id = str(row[0])
    source = primary_by_source.get(source_id)
    table2_consistent &= source is not None
    if source is not None:
        table2_consistent &= close(row[9], source["p"])
        table2_consistent &= close(row[10], source["q_primary_recomputed"])
        table2_consistent &= close(row[7], source["mean_F"])
add_check("table_2_matches_frozen_results", table2_consistent, "P, q and F checked for 7/7 rows")

figure3_rows = read_csv(FIGURES / "source_data" / "Figure_3a_source_data.csv")
figure3_consistent = len(figure3_rows) == 7
for row in figure3_rows:
    source = primary_by_source.get(row["source_id"])
    figure3_consistent &= source is not None
    if source is not None:
        figure3_consistent &= close(row["p"], source["p"])
        figure3_consistent &= close(row["q"], source["q_primary_recomputed"])
        figure3_consistent &= close(row["or"], source["or"])
add_check("figure_3_matches_frozen_results", figure3_consistent, "OR, P and q checked for 7/7 rows")

figure2_rows = read_csv(FIGURES / "source_data" / "Figure_2_source_data.csv")
forward_figure2 = [row for row in figure2_rows if row["direction"] == "Forward"]
reverse_figure2 = [row for row in figure2_rows if row["direction"] == "Reverse"]
add_check(
    "figure_2_complete_denominators",
    len(forward_figure2) == 230
    and len(reverse_figure2) == 1572
    and all(int(row["denominator"]) == 230 for row in forward_figure2)
    and all(int(row["denominator"]) == 1572 for row in reverse_figure2),
    f"forward={len(forward_figure2)}, reverse={len(reverse_figure2)}",
)
add_check(
    "figure_2_forward_nonestimable_retained",
    sum(close(row["observed_p"], 1.0) for row in forward_figure2) >= 12,
    f"forward P=1 rows={sum(close(row['observed_p'], 1.0) for row in forward_figure2)}",
)

supplement_wb = load_workbook(SUPPLEMENT / "IJIR_Supplementary_Data_v0_3.xlsx", read_only=True, data_only=True)
required_sheets = {
    "S1 Instruments",
    "S2 Forward Primary",
    "S3 Reverse Primary",
    "S4 Fwd Methods",
    "S4 Rev Methods",
    "S5 Power Summary",
    "S5 Power Traits",
    "S6 HUNT Same SNP",
    "S6 HUNT Nominal",
    "S7 Strict Results",
    "S8 Pleiotropy",
    "S9 Mechanistic Full",
    "S10 Provenance",
    "Closure Audit",
}
add_check(
    "supplementary_sheet_set",
    required_sheets.issubset(set(supplement_wb.sheetnames)),
    f"{len(supplement_wb.sheetnames)} sheets; required subset present={required_sheets.issubset(set(supplement_wb.sheetnames))}",
)
instrument_sheet_rows = sum(1 for _ in supplement_wb["S1 Instruments"].iter_rows())
add_check(
    "complete_instrument_inventory",
    instrument_sheet_rows >= 43298,
    f"worksheet rows including title/header={instrument_sheet_rows}",
)

closure = read_csv(RESULTS / "prespecified_analysis_closure_audit.csv")
closed_status = all(
    any(token in row["final_status"] for token in ("completed", "not_triggered"))
    for row in closure
)
add_check("methods_results_closure", len(closure) == 16 and closed_status, f"{len(closure)}/16 analyses closed")

hunt_summary = read_csv(RESULTS / "hunt_same_snp_validation_summary.csv")[0]
hunt_audit = read_csv(RESULTS / "hunt_nominal_trait_audit.csv")
add_check(
    "hunt_threshold_audit",
    len(hunt_audit) == 2
    and all(int(row["post_clump_snps_primary"]) == 0 for row in hunt_audit)
    and [int(row["post_clump_snps_exploratory"]) for row in hunt_audit] == [23, 20]
    and int(hunt_summary["same_position_found"]) == 97
    and int(hunt_summary["direction_concordant"]) == 76,
    "0 GWS instruments; exploratory sets 23/20; same-SNP coverage 97/97; concordance 76/97",
)

power = {row["alpha_label"]: row for row in read_csv(RESULTS / "power_mde_summary.csv")}
power_ok = (
    len(power) == 3
    and close(power["nominal_0.05"]["mde_or_median"], 2.17404162345676)
    and close(power["forward_bonferroni_0.05_over_230"]["mde_or_median"], 2.89351503171884)
    and close(power["global_bonferroni_0.05_over_1802"]["mde_or_median"], 3.09405706290943)
)
add_check("power_mde_recomputed", power_ok, "median MDE ORs 2.174, 2.894 and 3.094")

strict = read_csv(RESULTS / "source_study_wide_sensitivity_summary.csv")[0]
strict_ok = (
    int(strict["eligible_traits"]) == 27
    and int(strict["estimable_traits"]) == 23
    and int(strict["fdr_significant"]) == 0
    and close(strict["minimum_p"], 0.153256024848082)
)
add_check("source_study_wide_sensitivity", strict_ok, "27 eligible; 23 estimable; min P=0.153256; 0 FDR")

mechanistic = read_csv(RESULTS / "mechanistic_screening_family_summary.csv")
mechanistic_ok = (
    len(mechanistic) == 7
    and sum(int(row["planned_tests"]) for row in mechanistic) == 544
    and sum(int(row["non_estimable_tests"]) for row in mechanistic) == 185
    and sum(int(row["fdr_significant_tests"]) for row in mechanistic) == 0
)
add_check("mechanistic_workbook_complete", mechanistic_ok, "7 families; 544 rows; 185 non-estimable; 0 FDR")

pngs = [
    FIGURES / "IJIR_Figure_1_Study_Design_v0_3_600dpi.png",
    FIGURES / "IJIR_Figure_2_Multiplicity_v0_3_600dpi.png",
    FIGURES / "IJIR_Figure_3_Forward_Evidence_Profile_v0_3_600dpi.png",
    *sorted((FIGURES / "supplementary").glob("*_600dpi.png")),
]
dpi_evidence: list[str] = []
dpi_ok = len(pngs) == 8
for path in pngs:
    with Image.open(path) as image:
        dpi = image.info.get("dpi", (0, 0))
        dpi_evidence.append(f"{path.name}:{dpi[0]:.1f}")
        dpi_ok &= abs(float(dpi[0]) - 600.0) <= 1.0 and abs(float(dpi[1]) - 600.0) <= 1.0
add_check("figure_dpi", dpi_ok, "; ".join(dpi_evidence))

rendered_counts = {
    "main": len(list((QC / "rendered_main_final").glob("page-*.png"))),
    "title": len(list((QC / "rendered_title_final").glob("page-*.png"))),
    "cover": len(list((QC / "rendered_cover_final").glob("page-*.png"))),
    "strobe": len(list((QC / "rendered_strobe_final").glob("page-*.png"))),
}
add_check(
    "final_docx_render_receipts",
    rendered_counts == {"main": 16, "title": 1, "cover": 1, "strobe": 1},
    json.dumps(rendered_counts, sort_keys=True),
)

checksums = []
for path in required:
    if path.is_file():
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
        checksums.append({"file": str(path.relative_to(PACKAGE)), "bytes": path.stat().st_size, "sha256": digest})

summary = {
    "version": "v0.3",
    "generated": "2026-07-22",
    "passed": sum(bool(item["passed"]) for item in checks),
    "failed": sum(not bool(item["passed"]) for item in checks),
    "checks": checks,
    "checksums": checksums,
}
(QC / "FINAL_QA_MACHINE_CHECKS_v0_3.json").write_text(
    json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)
with (QC / "FINAL_DELIVERABLE_CHECKSUMS_v0_3.csv").open("w", newline="", encoding="utf-8") as handle:
    writer = csv.DictWriter(
        handle,
        fieldnames=["file", "bytes", "sha256"],
        lineterminator="\n",
    )
    writer.writeheader()
    writer.writerows(checksums)

for item in checks:
    status = "PASS" if item["passed"] else "FAIL"
    print(f"{status}\t{item['check']}\t{item['evidence']}")
print(f"SUMMARY\tpassed={summary['passed']}\tfailed={summary['failed']}")
raise SystemExit(0 if summary["failed"] == 0 else 1)
