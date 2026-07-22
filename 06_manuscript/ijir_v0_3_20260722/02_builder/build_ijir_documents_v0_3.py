#!/usr/bin/env python3
from pathlib import Path
import csv
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "01_sources"
OUT = ROOT / "03_submission_files"
QC = ROOT / "07_qc"
OUT.mkdir(parents=True, exist_ok=True)
QC.mkdir(parents=True, exist_ok=True)

MAIN_SOURCE = SRC / "IJIR_Main_Manuscript_v0_3.md"
TITLE_SOURCE = SRC / "IJIR_Title_Page_v0_3.md"
COVER_SOURCE = SRC / "IJIR_Cover_Letter_v0_3.md"
LOCATION_SOURCE = SRC / "STROBE_MR_locations_v0_3.csv"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def enable_line_numbers(section):
    sect_pr = section._sectPr
    line = sect_pr.find(qn("w:lnNumType"))
    if line is None:
        line = OxmlElement("w:lnNumType")
        sect_pr.append(line)
    line.set(qn("w:countBy"), "1")
    line.set(qn("w:start"), "1")
    line.set(qn("w:restart"), "continuous")
    line.set(qn("w:distance"), "360")


def prevent_widow_control(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    widow = OxmlElement("w:widowControl")
    widow.set(qn("w:val"), "0")
    p_pr.append(widow)


def setup_document(doc, *, double, line_numbers, margins=1.0):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(margins)
    section.left_margin = section.right_margin = Inches(margins)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = (
        WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.SINGLE
    )
    title = doc.styles["Title"]
    title.font.size = Pt(14)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(12)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    add_page_number(section)
    if line_numbers:
        enable_line_numbers(section)


def add_inline(paragraph, text):
    position = 0
    for match in re.finditer(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_markdown(doc, markdown_text, *, skip_title=False):
    lines = markdown_text.splitlines()
    pending = []

    def flush():
        nonlocal pending
        if pending:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, " ".join(line.strip() for line in pending))
            prevent_widow_control(paragraph)
            pending = []

    for line in lines:
        if not line.strip():
            flush()
            continue
        if line.startswith("# "):
            flush()
            if skip_title:
                continue
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            flush()
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            flush()
            doc.add_paragraph(line[4:], style="Heading 2")
        else:
            pending.append(line)
    flush()


def count_words(text):
    clean = re.sub(r"[#*]", "", text)
    return len(re.findall(r"\b[\w'’–-]+\b", clean, flags=re.UNICODE))


def extract_between(text, start_pattern, end_pattern):
    match = re.search(start_pattern + r"(.*?)" + end_pattern, text, re.S)
    if not match:
        raise RuntimeError(f"Could not extract text between {start_pattern} and {end_pattern}")
    return match.group(1).strip()


def read_locations():
    if not LOCATION_SOURCE.exists():
        return {}
    with LOCATION_SOURCE.open(newline="", encoding="utf-8") as handle:
        return {row["key"]: row["location"] for row in csv.DictReader(handle)}


main_md = MAIN_SOURCE.read_text(encoding="utf-8")
title_md = TITLE_SOURCE.read_text(encoding="utf-8")
cover_md = COVER_SOURCE.read_text(encoding="utf-8")

abstract = extract_between(main_md, r"## Abstract\n\n", r"\n\nKeywords:")
main_block = extract_between(main_md, r"## Introduction\n", r"\n\n## Data availability")
abstract_wc = count_words(abstract)
main_wc = count_words(main_block)
if abstract_wc > 200:
    raise RuntimeError(f"Abstract exceeds IJIR limit: {abstract_wc}")
if main_wc > 3000:
    raise RuntimeError(f"Main text exceeds IJIR limit: {main_wc}")
title_md = title_md.replace("[generated during build]", str(abstract_wc), 1)
title_md = title_md.replace("[generated during build]", str(main_wc), 1)


# Main manuscript: title and text, without duplicating the separate author page.
doc = Document()
setup_document(doc, double=True, line_numbers=True)
add_markdown(doc, main_md)
doc.core_properties.title = (
    "Bidirectional MR of gut microbial traits and erectile dysfunction"
)
doc.core_properties.subject = "IJIR Article manuscript v0.3"
doc.core_properties.author = "Xiancheng Du et al."
doc.core_properties.keywords = (
    "erectile dysfunction; gut microbiome; Mendelian randomization; HUNT"
)
doc.save(OUT / "IJIR_Main_Manuscript_v0_3.docx")


# Separate title page.
doc = Document()
setup_document(doc, double=False, line_numbers=True)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
add_markdown(doc, title_md)
doc.core_properties.title = "IJIR title page v0.3"
doc.core_properties.author = "Xiancheng Du et al."
doc.save(OUT / "IJIR_Title_Page_v0_3.docx")


# Cover letter.
doc = Document()
setup_document(doc, double=False, line_numbers=False, margins=0.65)
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(2)
add_markdown(doc, cover_md)
doc.core_properties.title = "Cover letter to International Journal of Impotence Research"
doc.core_properties.author = "Xiancheng Du"
doc.save(OUT / "IJIR_Cover_Letter_v0_3.docx")


locations = read_locations()


def loc(key, fallback):
    return locations.get(key, fallback)


checklist = [
    ("1", "Title and abstract", loc("title_abstract", "Title and Abstract"),
     "MR design, sources, thresholds, multiplicity, cross-cohort evidence and principal results are stated."),
    ("2", "Background", loc("background", "Introduction, paragraphs 1–2"),
     "Clinical rationale, prior microbiome evidence and reproducibility concerns are described."),
    ("3", "Objectives", loc("objectives", "Introduction, paragraph 3"),
     "Bidirectional multiplicity-controlled reassessment is specified."),
    ("4", "Study design and data sources", loc("design_sources", "Methods: Study design and data sources"),
     "Summary-data design, cohorts, ancestry, phenotype definitions and evidence roles are reported."),
    ("5", "Participants", loc("participants", "Methods: Study design and data sources; Table 1"),
     "Source-specific participant numbers and populations are reported."),
    ("6", "Genetic variants", loc("variants", "Methods: Instruments, harmonization, and effect scales"),
     "Selection thresholds, F statistic, clumping and LD reference are specified."),
    ("7", "Variables and measurement", loc("variables", "Methods: Study design and data sources; effect scales"),
     "Microbial and ED phenotypes and non-comparable scales are defined."),
    ("8", "MR assumptions", loc("assumptions", "Methods and Discussion"),
     "Relevance is enforced and independence/exclusion restrictions are addressed through design and limitations."),
    ("9", "Statistical methods", loc("statistics", "Methods: MR estimation, multiplicity, and diagnostic analyses"),
     "Wald ratio, IVW, robust estimators, complete BH families and sensitivity corrections are specified."),
    ("10", "Missing data and harmonization", loc("harmonization", "Methods: Instruments, harmonization, and effect scales"),
     "Genome build, allele, multiallelic, palindromic and non-estimable handling are described."),
    ("11", "Sensitivity analyses", loc("sensitivities", "Methods: source threshold, cross-cohort validation, power, and mechanistic screening"),
     "HUNT, ED ancestry, source-threshold, multiplicity, power and source-heterogeneity analyses are reported."),
    ("12", "Descriptive data", loc("descriptive", "Results: Primary bidirectional analyses; Table 1"),
     "Eligible, estimable, instrument-count and harmonization summaries are given."),
    ("13", "Main results", loc("main_results", "Results: Primary bidirectional analyses; Figures 2–3; Table 2"),
     "Effect estimates, P values, FDR values and denominators are reported without promoting nominal rows."),
    ("14", "Additional analyses", loc("additional", "Results: sensitivity, validation, detectability, and mechanistic screening"),
     "All prespecified analyses are completed, explicitly not triggered, or reported as non-estimable."),
    ("15", "Key results", loc("key_results", "Discussion, paragraph 1"),
     "The multiplicity-controlled conclusion is interpreted against the study objectives."),
    ("16", "Limitations", loc("limitations", "Discussion, paragraphs 3–7"),
     "Detectability, sparse instruments, pleiotropy, HUNT scale, ED phenotype, overlap and ancestry are addressed."),
    ("17", "Interpretation", loc("interpretation", "Discussion, paragraphs 1 and 8"),
     "Conclusions are bounded to tested instruments and do not claim proof of absence."),
    ("18", "Generalisability", loc("generalisability", "Discussion, paragraphs 5–8"),
     "Taxonomy, measurement, phenotype and ancestry transfer limitations are stated."),
    ("19", "Funding", loc("funding", "Funding"),
     "Grant numbers, recipients and funder roles are disclosed."),
    ("20", "Data, code and transparency", loc("transparency", "Data availability; Software and reproducibility"),
     "Source access, public archive status, provenance records and LLM use are disclosed."),
]


doc = Document()
setup_document(doc, double=False, line_numbers=False, margins=0.55)
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("STROBE-MR checklist").bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(
    "Bidirectional Mendelian randomization of gut microbial traits and erectile dysfunction "
    "with cross-cohort exposure validation and multiplicity control"
)
doc.add_paragraph(
    "Locations refer to the final rendered v0.3 main manuscript. The checklist "
    "records analyses that were completed, not triggered, or non-estimable."
)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.autofit = False
headers = ["Item", "Reporting domain", "Final manuscript page/line", "Implementation note"]
for index, heading in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.text = heading
    set_cell_shading(cell, "52677D")
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        run.font.size = Pt(8)
set_repeat_header(table.rows[0])
for row_index, row in enumerate(checklist, start=1):
    cells = table.add_row().cells
    if row_index % 2 == 0:
        for cell in cells:
            set_cell_shading(cell, "F3F6F8")
    for index, value in enumerate(row):
        cells[index].text = value
        cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for paragraph in cells[index].paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
widths = [0.45, 1.65, 2.40, 5.00]
for row in table.rows:
    for index, width in enumerate(widths):
        set_cell_width(row.cells[index], width)
doc.core_properties.title = "STROBE-MR checklist v0.3"
doc.core_properties.author = "Xiancheng Du et al."
doc.save(OUT / "IJIR_STROBE_MR_Checklist_v0_3.docx")


word_report = QC / "IJIR_Word_Counts_v0_3.txt"
word_report.write_text(
    f"Abstract words: {abstract_wc}\n"
    f"Main-text words (Introduction through Discussion): {main_wc}\n"
    f"References: {len(re.findall(r'^\d+\.', main_md, flags=re.M))}\n"
    f"Main figures: 3\nMain tables: 2\nSupplementary files: 3\n",
    encoding="utf-8",
)
print(
    f"Created v0.3 DOCX files. Abstract words={abstract_wc}; "
    f"main-text words={main_wc}."
)
