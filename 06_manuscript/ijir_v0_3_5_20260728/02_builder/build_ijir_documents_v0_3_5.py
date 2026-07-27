#!/usr/bin/env python3
from pathlib import Path
import csv
import json
import re

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "01_sources"
OUT = ROOT / "03_submission_files"
QC = ROOT / "07_qc"
OUT.mkdir(parents=True, exist_ok=True)
QC.mkdir(parents=True, exist_ok=True)

MAIN_SOURCE = SRC / "IJIR_Main_Manuscript_v0_3_5.md"
TITLE_SOURCE = SRC / "IJIR_Title_Page_v0_3_5.md"
COVER_SOURCE = SRC / "IJIR_Cover_Letter_v0_3_5.md"
LOCATION_SOURCE = SRC / "STROBE_MR_locations_v0_3_5.csv"
TABLE_DATA = ROOT / "02_builder" / "tmp_table_data"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def enable_line_numbers(section):
    sect_pr = section._sectPr
    line = sect_pr.find(qn("w:lnNumType"))
    if line is None:
        line = OxmlElement("w:lnNumType")
        sect_pr.append(line)
    line.set(qn("w:countBy"), "1")
    line.set(qn("w:start"), "1")
    line.set(qn("w:restart"), "continuous")
    line.set(qn("w:distance"), "360")


def prevent_widow_control(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    widow = OxmlElement("w:widowControl")
    widow.set(qn("w:val"), "0")
    p_pr.append(widow)


def setup_document(doc, *, double, line_numbers, margins=1.0):
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(margins)
    section.left_margin = section.right_margin = Inches(margins)
    for name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        p_pr = style._element.get_or_add_pPr()
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is not None:
            p_pr.remove(p_bdr)
    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing_rule = (
        WD_LINE_SPACING.DOUBLE if double else WD_LINE_SPACING.SINGLE
    )
    title = doc.styles["Title"]
    title.font.size = Pt(14)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(12)
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        style = doc.styles[name]
        style.font.size = Pt(12)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.keep_with_next = True
    add_page_number(section)
    if line_numbers:
        enable_line_numbers(section)


def set_core_metadata(doc, *, title, author):
    props = doc.core_properties
    props.title = title
    props.subject = "International Journal of Impotence Research submission"
    props.author = author
    props.last_modified_by = "Xiancheng Du"
    props.category = "Article submission"
    props.version = "v0.3.5"


def add_inline(paragraph, text):
    position = 0
    for match in re.finditer(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_markdown(doc, markdown_text, *, skip_title=False):
    lines = markdown_text.splitlines()
    pending = []

    def flush():
        nonlocal pending
        if pending:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, " ".join(line.strip() for line in pending))
            prevent_widow_control(paragraph)
            pending = []

    for line in lines:
        if not line.strip():
            flush()
            continue
        if line.startswith("# "):
            flush()
            if skip_title:
                continue
            paragraph = doc.add_paragraph(style="Title")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, line[2:])
        elif line.startswith("## "):
            flush()
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("### "):
            flush()
            doc.add_paragraph(line[4:], style="Heading 2")
        else:
            pending.append(line)
    flush()


def set_cell_margins(cell, top=40, start=50, bottom=40, end=50):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def compact_cell(cell, value, *, size=7.2, bold=False, white=False):
    cell.text = str(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        for run in paragraph.runs:
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0)


def add_compact_landscape_table(doc, *, title, headers, rows, widths, footnote):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = section.bottom_margin = Inches(0.38)
    section.left_margin = section.right_margin = Inches(0.40)

    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(5)
    run = heading.add_run(title)
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)

    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, header in enumerate(headers):
        compact_cell(table.rows[0].cells[index], header, size=7.0, bold=True, white=True)
        set_cell_shading(table.rows[0].cells[index], "3E5267")
        set_cell_width(table.rows[0].cells[index], widths[index])
    set_repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])

    for row_index, values in enumerate(rows, start=1):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            compact_cell(row.cells[index], value, size=7.0)
            set_cell_width(row.cells[index], widths[index])
            if row_index % 2 == 0:
                set_cell_shading(row.cells[index], "F2F4F6")

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.line_spacing = 1.0
    run = note.add_run(footnote)
    run.font.name = "Times New Roman"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor(0, 0, 0)


def sci(value):
    if value in ("", None):
        return ""
    value = float(value)
    if value == 0:
        return "0"
    if 0.001 <= abs(value) < 1000:
        return f"{value:.4f}".rstrip("0").rstrip(".")
    exponent = int(f"{value:.3e}".split("e")[1])
    mantissa = float(f"{value:.3e}".split("e")[0])
    return f"{mantissa:.3g}×10^{exponent}"


def build_main_table_rows():
    with (TABLE_DATA / "main_table_1.json").open(encoding="utf-8") as handle:
        table1 = json.load(handle)
    with (TABLE_DATA / "main_table_2.json").open(encoding="utf-8") as handle:
        table2 = json.load(handle)

    rows1 = [
        [
            row["GWAS source"],
            f'{row["Ancestry"]}; {row["Sample size"]}',
            row["Phenotype and effect scale"].replace(
                "Z/sqrt(weight)", "Z/√weight"
            ),
            row["Analytical role"],
            row["Overlap and evidence classification"],
        ]
        for row in table1
    ]

    rows2 = []
    for row in table2:
        trait = row["Microbial trait"].replace("Gut microbiome ", "", 1)
        phenotype = row["Exposure phenotype"]
        scale = (
            "Presence; OR per 1-unit genetically predicted log odds"
            if phenotype == "presence"
            else "RIN abundance; OR per 1 standardized unit"
        )
        if row["Exact HUNT match"] == "Yes":
            if "rs56024701" in row["Swedish lead SNP"]:
                hunt = "Exact label; same-SNP P=0.965, F=0.002; exploratory 23-SNP MR P=0.427"
            else:
                hunt = "Exact label; same-SNP P=0.671, F=0.180; exploratory 20-SNP MR P=0.175"
        else:
            hunt = "No exact label; no fuzzy matching"
        rows2.append(
            [
                trait,
                scale,
                f'{row["Swedish lead SNP"]}; F={float(row["F statistic"]):.2f}',
                row["Discovery OR (95% CI)"],
                sci(row["Nominal P"]),
                f'{float(row["FDR q"]):.4f}',
                hunt,
                row["Taxonomic signal cluster"],
            ]
        )
    return rows1, rows2


def count_words(text):
    clean = re.sub(r"[#*]", "", text)
    return len(re.findall(r"\b[\w'’–-]+\b", clean, flags=re.UNICODE))


def extract_between(text, start_pattern, end_pattern):
    match = re.search(start_pattern + r"(.*?)" + end_pattern, text, re.S)
    if not match:
        raise RuntimeError(f"Could not extract text between {start_pattern} and {end_pattern}")
    return match.group(1).strip()


def read_locations():
    if not LOCATION_SOURCE.exists():
        return {}
    with LOCATION_SOURCE.open(newline="", encoding="utf-8") as handle:
        return {row["key"]: row["location"] for row in csv.DictReader(handle)}


main_md = MAIN_SOURCE.read_text(encoding="utf-8")
title_md = TITLE_SOURCE.read_text(encoding="utf-8")
cover_md = COVER_SOURCE.read_text(encoding="utf-8")

abstract = extract_between(main_md, r"## Abstract\n\n", r"\n\nKeywords:")
main_block = extract_between(main_md, r"## Introduction\n", r"\n\n## Data availability")
abstract_wc = count_words(abstract)
main_wc = count_words(main_block)
if abstract_wc > 200:
    raise RuntimeError(f"Abstract exceeds IJIR limit: {abstract_wc}")
if main_wc > 3000:
    raise RuntimeError(f"Main text exceeds IJIR limit: {main_wc}")
title_md = title_md.replace("[generated during build]", str(abstract_wc), 1)
title_md = title_md.replace("[generated during build]", str(main_wc), 1)


# Main manuscript: title and text, without duplicating the separate author page.
doc = Document()
setup_document(doc, double=True, line_numbers=True)
main_without_table_notes = main_md.split("\n\n## Table titles and footnotes", 1)[0]
add_markdown(doc, main_without_table_notes)
table1_rows, table2_rows = build_main_table_rows()
add_compact_landscape_table(
    doc,
    title="Table 1. GWAS sources and analytical roles.",
    headers=[
        "GWAS source",
        "Ancestry and sample size",
        "Phenotype/effect scale",
        "Analytical role",
        "Overlap/evidence classification",
    ],
    rows=table1_rows,
    widths=[1.70, 1.65, 2.70, 1.85, 2.30],
    footnote=(
        "The 2025 European and cross-ancestry analyses include FinnGen; the "
        "African-ancestry stratum does not. HUNT is an independent exposure "
        "cohort, but FinnGen remains the MR outcome;\nNeither the legacy "
        "exploratory analysis nor the same-SNP lookup constitutes independent "
        "MR replication."
    ),
)
add_compact_landscape_table(
    doc,
    title="Table 2. Nominal forward gut microbial trait–ED associations.",
    headers=[
        "Microbial trait",
        "Exposure phenotype/scale",
        "Lead SNP and F",
        "OR (95% CI)",
        "Nominal P",
        "FDR q",
        "Concise HUNT evaluation",
        "Taxonomic/signal cluster",
    ],
    rows=table2_rows,
    widths=[1.85, 1.45, 0.95, 1.12, 0.58, 0.55, 2.00, 1.70],
    footnote=(
        "These are seven nominal trait-level associations, not seven independent "
        "causal taxa; none survived FDR correction. Peptococcaceae, "
        "Peptococcales, and Peptococcia are nested traits with identical "
        "estimates. Swedish presence and HUNT normalized relative-abundance "
        "scales are not directly comparable. The two focal exact-label HUNT "
        "GWASs had no genome-wide-significant instruments; legacy 23-SNP and "
        "20-SNP analyses used P<1×10⁻⁵ and were exploratory. All Swedish "
        "nominal estimates were single-SNP Wald ratios."
    ),
)
set_core_metadata(
    doc,
    title="Bidirectional MR of gut microbial traits and erectile dysfunction",
    author="Xiancheng Du et al.",
)
doc.core_properties.keywords = (
    "erectile dysfunction; gut microbiome; Mendelian randomization; HUNT"
)
doc.save(OUT / "IJIR_Main_Manuscript_v0_3_5.docx")


# Separate title page.
doc = Document()
setup_document(doc, double=False, line_numbers=True)
doc.styles["Normal"].paragraph_format.line_spacing = 1.15
add_markdown(doc, title_md)
set_core_metadata(doc, title="IJIR title page", author="Xiancheng Du et al.")
doc.save(OUT / "IJIR_Title_Page_v0_3_5.docx")


# Cover letter.
doc = Document()
setup_document(doc, double=False, line_numbers=False, margins=0.65)
doc.styles["Normal"].font.size = Pt(10.5)
doc.styles["Normal"].paragraph_format.space_after = Pt(2)
add_markdown(doc, cover_md)
set_core_metadata(
    doc,
    title="Cover letter to International Journal of Impotence Research",
    author="Xiancheng Du",
)
doc.save(OUT / "IJIR_Cover_Letter_v0_3_5.docx")


locations = read_locations()


def loc(key, fallback):
    return locations.get(key, fallback)


checklist = [
    ("1", "Title and abstract", loc("title_abstract", "Title and Abstract"),
     "MR design, sources, thresholds, multiplicity, cross-cohort evidence and principal results are stated."),
    ("2", "Background", loc("background", "Introduction, paragraphs 1–2"),
     "Clinical rationale, prior microbiome evidence and reproducibility concerns are described."),
    ("3", "Objectives", loc("objectives", "Introduction, paragraph 3"),
     "Bidirectional multiplicity-controlled reassessment is specified."),
    ("4", "Study design and data sources", loc("design_sources", "Methods: Study design and data sources"),
     "Summary-data design, cohorts, ancestry, phenotype definitions and evidence roles are reported."),
    ("5", "Participants", loc("participants", "Methods: Study design and data sources; Table 1"),
     "Source-specific participant numbers and populations are reported."),
    ("6", "Genetic variants", loc("variants", "Methods: Instruments, harmonization, and effect scales"),
     "Selection thresholds, F statistic, clumping and LD reference are specified."),
    ("7", "Variables and measurement", loc("variables", "Methods: Study design and data sources; effect scales"),
     "Microbial and ED phenotypes and non-comparable scales are defined."),
    ("8", "MR assumptions", loc("assumptions", "Methods and Discussion"),
     "Relevance is enforced and independence/exclusion restrictions are addressed through design and limitations."),
    ("9", "Statistical methods", loc("statistics", "Methods: MR estimation, multiplicity, and diagnostic analyses"),
     "Wald ratio, IVW, robust estimators, complete BH families and sensitivity corrections are specified."),
    ("10", "Missing data and harmonization", loc("harmonization", "Methods: Instruments, harmonization, and effect scales"),
     "Genome build, allele, multiallelic, palindromic and non-estimable handling are described."),
    ("11", "Sensitivity analyses", loc("sensitivities", "Methods: source threshold, cross-cohort evaluation, power, and mechanistic screening"),
     "HUNT, alternative ED outcomes, ancestry transfer, source-threshold, multiplicity, power and source-heterogeneity analyses defined before screening are separated from the targeted post hoc locus audit."),
    ("12", "Descriptive data", loc("descriptive", "Results: Primary bidirectional analyses; Table 1"),
     "Eligible, estimable, instrument-count and harmonization summaries are given."),
    ("13", "Main results", loc("main_results", "Results: Primary bidirectional analyses; Figures 2–3; Table 2"),
     "Effect estimates, P values, FDR values and denominators are reported without promoting nominal rows."),
    ("14", "Additional analyses", loc("additional", "Results: outcome sensitivity, same-SNP evaluation, detectability, and mechanistic screening"),
     "All analyses defined before screening are completed, explicitly not triggered, or non-estimable; each alternative outcome uses a separate 230-trait BH family, and the explanatory locus audit is labelled post hoc."),
    ("15", "Key results", loc("key_results", "Discussion, paragraph 1"),
     "The multiplicity-controlled conclusion is interpreted against the study objectives."),
    ("16", "Limitations", loc("limitations", "Discussion, paragraphs 3–7"),
     "Detectability, sparse instruments, locus redundancy, pleiotropy, HUNT scale, ED phenotype, overlap and ancestry transportability are addressed."),
    ("17", "Interpretation", loc("interpretation", "Discussion, paragraphs 1 and 8"),
     "Outcome discordance is acknowledged; the post hoc audit is explanatory rather than part of the original multiplicity criteria, and no proof of absence is claimed."),
    ("18", "Generalisability", loc("generalisability", "Discussion, paragraphs 5–8"),
     "Taxonomy, measurement, phenotype and ancestry transfer limitations are stated."),
    ("19", "Funding", loc("funding", "Funding"),
     "Grant numbers, recipients and funder roles are disclosed."),
    ("20", "Data, code and transparency", loc("transparency", "Data availability; Software and reproducibility"),
     "Source access, public archive status, provenance records and LLM use are disclosed."),
]


doc = Document()
setup_document(doc, double=False, line_numbers=False, margins=0.55)
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width, section.page_height = section.page_height, section.page_width
p = doc.add_paragraph(style="Title")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("STROBE-MR checklist").bold = True
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run(
    "Bidirectional Mendelian randomization of gut microbial traits and erectile dysfunction "
    "with cross-cohort same-SNP evaluation and multiplicity control"
)
doc.add_paragraph(
    "Locations refer to the final rendered main manuscript. The checklist "
    "records analyses that were completed, not triggered, or non-estimable."
)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
table.autofit = False
headers = ["Item", "Reporting domain", "Final manuscript page/line", "Implementation note"]
for index, heading in enumerate(headers):
    cell = table.rows[0].cells[index]
    cell.text = heading
    set_cell_shading(cell, "52677D")
    for run in cell.paragraphs[0].runs:
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.bold = True
        run.font.size = Pt(8)
set_repeat_header(table.rows[0])
for row_index, row in enumerate(checklist, start=1):
    cells = table.add_row().cells
    if row_index % 2 == 0:
        for cell in cells:
            set_cell_shading(cell, "F3F6F8")
    for index, value in enumerate(row):
        cells[index].text = value
        cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for paragraph in cells[index].paragraphs:
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            for run in paragraph.runs:
                run.font.name = "Times New Roman"
                run.font.size = Pt(8)
widths = [0.45, 1.65, 2.40, 5.00]
for row in table.rows:
    for index, width in enumerate(widths):
        set_cell_width(row.cells[index], width)
set_core_metadata(
    doc, title="STROBE-MR checklist", author="Xiancheng Du et al."
)
doc.save(OUT / "IJIR_STROBE_MR_Checklist_v0_3_5.docx")


word_report = QC / "IJIR_Word_Counts_v0_3_5.txt"
word_report.write_text(
    f"Abstract words: {abstract_wc}\n"
    f"Main-text words (Introduction through Discussion): {main_wc}\n"
    f"References: {len(re.findall(r'^\d+\.', main_md, flags=re.M))}\n"
    f"Main figures: 3\nMain tables: 2\nSupplementary files: 3\n",
    encoding="utf-8",
)
print(
    f"Created v0.3.5 DOCX files. Abstract words={abstract_wc}; "
    f"main-text words={main_wc}."
)
